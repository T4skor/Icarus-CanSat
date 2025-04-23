import serial
import matplotlib.pyplot as plt
import matplotlib.animation as animation
import pandas as pd
import re
import numpy as np
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import Series
from openpyxl.utils.dataframe import dataframe_to_rows  # Importa la función necesaria

# Configuración del puerto serie
puerto = 'COM6'  # Cambia al puerto correspondiente
baudrate = 9600  # Debe coincidir con la configuración de Arduino

try:
    ser = serial.Serial(puerto, baudrate)
    print("Conexión establecida con el puerto serie.")
except Exception as e:
    print(f"Error al conectar con el puerto serie: {e}")
    exit()

# Listas para los gráficos (si deseas visualizar la evolución)
tiempo_graph = []
temperaturas_graph = []
altitudes_graph = []
magX_graph = []
magY_graph = []
magZ_graph = []

# Expresiones regulares para extraer datos
pattern_bmp = r"Icarus:\s*Temperatura:\s*([-+]?\d+\.\d+)\s*°C,\s*Presión:\s*([-+]?\d+\.\d+)\s*atm,\s*Altitud:\s*([-+]?\d+\.\d+)\s*m"
pattern_mag = r"Icarus:\s*X:\s*(-?\d+\.\d+|\d+)\s*µT\s*Icarus:\s*Y:\s*(-?\d+\.\d+|\d+)\s*µT\s*Icarus:\s*Z:\s*(-?\d+\.\d+|\d+)"

# DataFrame global para guardar todos los datos de los sensores
df = pd.DataFrame(columns=["Tiempo", "Sensor", "Temperatura", "Presión", "Altitud", "MagX", "MagY", "MagZ"])

# Configuración de la figura con 4 subgráficas
fig, (ax1, ax2, ax3, ax4) = plt.subplots(4, 1, figsize=(10, 10))

# Gráfica 1: Temperatura vs Tiempo
ax1.set_title("Temperatura vs Tiempo")
ax1.set_xlabel("Tiempo (s)")
ax1.set_ylabel("Temperatura (°C)")

# Gráfica 2: Altitud vs Tiempo
ax2.set_title("Altitud vs Tiempo")
ax2.set_xlabel("Tiempo (s)")
ax2.set_ylabel("Altitud (m)")

# Gráfica 3: Gráfica Existente (por ejemplo, Temperatura vs Tiempo)
ax3.set_title("Gráfica Existente")
ax3.set_xlabel("Tiempo (s)")
ax3.set_ylabel("Temperatura (°C)")

# Gráfica 4: Magnetómetro (X, Y, Z)
ax4.set_title("Magnetómetro (X, Y, Z)")
ax4.set_xlabel("Tiempo (s)")
ax4.set_ylabel("Magnetómetro (µT)")
line_x, = ax4.plot([], [], label="X", color='r')
line_y, = ax4.plot([], [], label="Y", color='g')
line_z, = ax4.plot([], [], label="Z", color='b')
ax4.legend()

def actualizar(i):
    global df
    # Revisar si hay datos en el buffer del puerto serie
    if ser.in_waiting > 0:
        linea = ser.readline().decode('utf-8').strip()
        print(f"Línea recibida: {linea}")
        
        # Usamos i (índice de actualización) como marcador de tiempo
        # Idealmente podrías usar time.time() o similar para la marca real
        tiempo_actual = i
        
        # Verificar si la línea corresponde a datos del BMP280
        if "Temperatura:" in linea:
            m_bmp = re.search(pattern_bmp, linea)
            if m_bmp:
                temp_val = float(m_bmp.group(1))
                pres_val = float(m_bmp.group(2))
                alt_val = float(m_bmp.group(3))
                
                # Actualizar listas para gráficos
                tiempo_graph.append(tiempo_actual)
                temperaturas_graph.append(temp_val)
                altitudes_graph.append(alt_val)
                
                # Crear una fila para los datos del BMP280
                new_row = {
                    "Tiempo": tiempo_actual,
                    "Sensor": "BMP280",
                    "Temperatura": temp_val,
                    "Presión": pres_val,
                    "Altitud": alt_val,
                    "MagX": None,
                    "MagY": None,
                    "MagZ": None
                }
                df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
                
        # Verificar si la línea corresponde a datos del magnetómetro
        elif "X:" in linea:
            m_mag = re.search(pattern_mag, linea)
            if m_mag:
                x_val = float(m_mag.group(1))
                y_val = float(m_mag.group(2))
                z_val = float(m_mag.group(3))
                
                # Actualizar listas para la gráfica del magnetómetro
                magX_graph.append(x_val)
                magY_graph.append(y_val)
                magZ_graph.append(z_val)
                if len(magX_graph) > 100:
                    magX_graph.pop(0)
                    magY_graph.pop(0)
                    magZ_graph.pop(0)
                
                # Crear una fila para los datos del magnetómetro
                new_row = {
                    "Tiempo": tiempo_actual,
                    "Sensor": "Magnetómetro",
                    "Temperatura": None,
                    "Presión": None,
                    "Altitud": None,
                    "MagX": x_val,
                    "MagY": y_val,
                    "MagZ": z_val
                }
                df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
                
                # Actualizar la gráfica del magnetómetro (ax4)
                line_x.set_data(range(len(magX_graph)), magX_graph)
                line_y.set_data(range(len(magY_graph)), magY_graph)
                line_z.set_data(range(len(magZ_graph)), magZ_graph)
                ax4.relim()
                ax4.autoscale_view()
        
        # Actualizar las gráficas usando los datos almacenados
        # Gráfica 1: Temperatura vs Tiempo
        ax1.clear()
        ax1.plot(tiempo_graph, temperaturas_graph, color='red')
        ax1.set_title("Temperatura vs Tiempo")
        ax1.set_xlabel("Tiempo (s)")
        ax1.set_ylabel("Temperatura (°C)")
        
        # Gráfica 2: Altitud vs Tiempo
        ax2.clear()
        ax2.plot(tiempo_graph, altitudes_graph, color='blue')
        ax2.set_title("Altitud vs Tiempo")
        ax2.set_xlabel("Tiempo (s)")
        ax2.set_ylabel("Altitud (m)")
        
        # Gráfica 3: Gráfica Existente (por ejemplo, Temperatura vs Tiempo)
        ax3.clear()
        ax3.plot(tiempo_graph, temperaturas_graph, label="Temperatura", color="green")
        ax3.set_title("Tiempo vs Temperatura")
        ax3.set_xlabel("Tiempo (s)")
        ax3.set_ylabel("Temperatura (°C)")
        ax3.legend()
        
        # Guardar el DataFrame completo en un archivo Excel
        try:
            df.to_excel("datos_sensores.xlsx", index=False)

            # Crear el archivo Excel y agregar los gráficos
            wb = Workbook()
            ws = wb.active
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)

            # Gráfico de Temperatura vs Tiempo
            chart1 = LineChart()
            data1 = Reference(ws, min_col=3, min_row=1, max_col=3, max_row=len(df)+1)
            categories1 = Reference(ws, min_col=1, min_row=2, max_row=len(df)+1)
            chart1.add_data(data1, titles_from_data=True)
            chart1.set_categories(categories1)
            chart1.title = "Temperatura vs Tiempo"
            chart1.x_axis.title = "Tiempo (s)"
            chart1.y_axis.title = "Temperatura (°C)"
            chart1.series[0].graphicalProperties.line.solidFill = "FF0000"  # Color rojo
            ws.add_chart(chart1, "E5")

            # Gráfico de Altitud vs Tiempo
            chart2 = LineChart()
            data2 = Reference(ws, min_col=4, min_row=1, max_col=4, max_row=len(df)+1)
            categories2 = Reference(ws, min_col=1, min_row=2, max_row=len(df)+1)
            chart2.add_data(data2, titles_from_data=True)
            chart2.set_categories(categories2)
            chart2.title = "Altitud vs Tiempo"
            chart2.x_axis.title = "Tiempo (s)"
            chart2.y_axis.title = "Altitud (m)"
            chart2.series[0].graphicalProperties.line.solidFill = "0000FF"  # Color azul
            ws.add_chart(chart2, "E20")

            # Gráfico de Magnetómetro X, Y, Z vs Tiempo
            chart3 = LineChart()
            data3 = Reference(ws, min_col=6, min_row=1, max_col=8, max_row=len(df)+1)
            categories3 = Reference(ws, min_col=1, min_row=2, max_row=len(df)+1)
            chart3.add_data(data3, titles_from_data=True)
            chart3.set_categories(categories3)
            chart3.title = "Magnetómetro X, Y, Z vs Tiempo"
            chart3.x_axis.title = "Tiempo (s)"
            chart3.y_axis.title = "µT"
            chart3.series[0].graphicalProperties.line.solidFill = "FF0000"  # Color rojo para X
            chart3.series[1].graphicalProperties.line.solidFill = "00FF00"  # Color verde para Y
            chart3.series[2].graphicalProperties.line.solidFill = "0000FF"  # Color azul para Z
            ws.add_chart(chart3, "E35")

            wb.save("datos_sensores_con_graficas.xlsx")
            print("Datos guardados en el archivo Excel con gráficos.")

        except Exception as e:
            print(f"Error al guardar el archivo Excel: {e}")
        
        plt.draw()
        plt.pause(0.001)

ani = animation.FuncAnimation(fig, actualizar, interval=100, save_count=200, cache_frame_data=False)

plt.show()
ser.close()
